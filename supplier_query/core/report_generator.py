# -*- coding: utf-8 -*-
"""
报告生成模块 - 生成Word文档报告
"""

from pathlib import Path
from datetime import datetime
from typing import List

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from core.query_engine import SiteResult, QueryStatus
from utils.logger import logger
from utils.file_utils import file_utils


class ReportGenerator:
    """Word报告生成器"""

    def __init__(self, output_dir: Path):
        """
        初始化报告生成器

        Args:
            output_dir: 输出目录
        """
        self.output_dir = output_dir

    def generate(
        self,
        supplier_name: str,
        credit_code: str,
        is_engineering: bool,
        results: List[SiteResult]
    ) -> Path:
        """
        生成Word报告

        Args:
            supplier_name: 供应商名称
            credit_code: 统一社会信用代码
            is_engineering: 是否工程类
            results: 查询结果列表

        Returns:
            报告文件路径
        """
        try:
            # 创建文档
            doc = Document()

            # 设置文档样式
            self._setup_document_style(doc)

            # 添加标题
            self._add_title(doc)

            # 添加基本信息表
            self._add_basic_info(doc, supplier_name, credit_code, is_engineering)

            # 添加查询结果截图
            self._add_screenshots(doc, results)

            # 保存文档
            filename = file_utils.generate_report_name(supplier_name)
            save_path = self.output_dir / filename
            doc.save(str(save_path))

            logger.info(f"报告已生成: {save_path}")
            return save_path

        except Exception as e:
            logger.error(f"报告生成失败: {e}")
            raise

    def _setup_document_style(self, doc: Document):
        """设置文档样式"""
        # 设置默认字体（跨平台支持中文）
        style = doc.styles['Normal']
        style.font.name = 'Microsoft YaHei'
        style._element.rPr.rFonts.set(
            doc.styles['Normal']._element.rPr.rFonts.attrib,
            'eastAsia',
            'Microsoft YaHei'
        )
        style.font.size = Pt(10.5)

    def _add_title(self, doc: Document):
        """添加标题"""
        title = doc.add_heading('供应商不良行为查询报告', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 设置标题字体
        for run in title.runs:
            run.font.name = 'Microsoft YaHei'
            run._element.rPr.rFonts.set(
                run._element.rPr.rFonts.attrib,
                'eastAsia',
                'Microsoft YaHei'
            )
            run.font.size = Pt(22)
            run.font.bold = True

    def _add_basic_info(self, doc: Document, supplier_name: str, credit_code: str, is_engineering: bool):
        """添加基本信息表"""
        doc.add_paragraph()  # 空行

        # 创建表格
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Table Grid'

        # 设置列宽
        for row in table.rows:
            row.cells[0].width = Cm(4)
            row.cells[1].width = Cm(10)

        # 填充数据
        info_data = [
            ("供应商名称", supplier_name),
            ("统一社会信用代码", credit_code),
            ("是否工程类", "是" if is_engineering else "否"),
            ("查询时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("查询日期", datetime.now().strftime("%Y-%m-%d")),
        ]

        for i, (label, value) in enumerate(info_data):
            row = table.rows[i]
            row.cells[0].text = label
            row.cells[1].text = value

            # 格式化
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.name = 'Microsoft YaHei'
                        run._element.rPr.rFonts.set(
                            run._element.rPr.rFonts.attrib,
                            'eastAsia',
                            'Microsoft YaHei'
                        )
                        run.font.size = Pt(10.5)

        doc.add_paragraph()  # 空行

    def _add_screenshots(self, doc: Document, results: List[SiteResult]):
        """添加截图到文档"""
        doc.add_heading('查询结果截图', level=1)

        success_results = [r for r in results if r.status == QueryStatus.SUCCESS and r.screenshot_path]

        if not success_results:
            doc.add_paragraph("（未获取到任何截图）")
            return

        for result in success_results:
            # 添加网站名称标签
            site_title = doc.add_heading(result.site_name, level=2)

            # 设置标题样式
            for run in site_title.runs:
                run.font.name = 'Microsoft YaHei'
                run._element.rPr.rFonts.set(
                    run._element.rPr.rFonts.attrib,
                    'eastAsia',
                    'Microsoft YaHei'
                )
                run.font.size = Pt(14)
                run.font.bold = True

            # 添加截图
            if result.screenshot_path and result.screenshot_path.exists():
                try:
                    # 设置图片宽度（保持比例）
                    doc.add_picture(
                        str(result.screenshot_path),
                        width=Inches(5.5)
                    )

                    # 图片居中
                    last_paragraph = doc.paragraphs[-1]
                    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                except Exception as e:
                    logger.warning(f"添加图片失败: {result.screenshot_path}, 错误: {e}")
                    doc.add_paragraph(f"（图片加载失败: {result.screenshot_path.name}）")
            else:
                doc.add_paragraph("（截图不存在）")

            # 空行
            doc.add_paragraph()

        # 添加查询统计
        self._add_summary(doc, results)

    def _add_summary(self, doc: Document, results: List[SiteResult]):
        """添加查询统计"""
        doc.add_heading('查询统计', level=1)

        total = len(results)
        success = sum(1 for r in results if r.status == QueryStatus.SUCCESS)
        failed = sum(1 for r in results if r.status == QueryStatus.FAILED)

        summary_text = f"本次共查询 {total} 个站点，其中成功 {success} 个，失败 {failed} 个。"

        para = doc.add_paragraph(summary_text)
        for run in para.runs:
            run.font.name = 'Microsoft YaHei'
            run._element.rPr.rFonts.set(
                run._element.rPr.rFonts.attrib,
                'eastAsia',
                'Microsoft YaHei'
            )

        # 失败站点列表
        if failed > 0:
            doc.add_paragraph()
            doc.add_heading('失败站点', level=2)
            for r in results:
                if r.status == QueryStatus.FAILED:
                    doc.add_paragraph(f"• {r.site_name}: {r.error_message}")


# 全局实例工厂
def create_report_generator(output_dir: Path) -> ReportGenerator:
    """创建报告生成器实例"""
    return ReportGenerator(output_dir)